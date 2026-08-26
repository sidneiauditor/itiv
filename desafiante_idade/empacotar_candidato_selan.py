# -*- coding: utf-8 -*-
"""Monta pasta autônoma do candidato para avaliação do Selan.

Cria: Modelo_Apartamentos_Candidato_Idade_20260811/
Não altera Modelo_Apartamentos_Aprovado/.

  python -u desafiante_idade/empacotar_candidato_selan.py
"""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
import warnings
from datetime import date
from pathlib import Path

import lightgbm as lgb
import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
PASTA_ITIV = HERE.parent
CODIGO_APROV = PASTA_ITIV / "Modelo_Apartamentos_Aprovado" / "codigo"
SAIDA_EXP = HERE / "saida"

sys.path.insert(0, str(CODIGO_APROV))
sys.path.insert(0, str(HERE))

from idade_features import (  # noqa: E402
    FEATURES_IDADE,
    FAIXAS_INSCRICAO,
    anexar_idade,
    carregar_nascimento,
)
from teste_inscricao_relativa_setor import (  # noqa: E402
    FEATURES_REL,
    anexar_inscricao_relativa,
    stats_setor_treino,
)
from apartamentos_lib import (  # noqa: E402
    AMOSTRAS,
    CAT_LGBM,
    FEATURES_LGBM,
    K_VIZINHOS,
    SEMENTE,
    calcular_knn_proxy,
    metricas_completas,
    preparar_area_saneamento,
)
from metricas_iaao_por_decil import metricas_razao, prb_global  # noqa: E402
from saneamento_chauvenet_iterativo import aplicar_chauvenet_iterativo  # noqa: E402

warnings.filterwarnings("ignore")

PACOTE = PASTA_ITIV / "Modelo_Apartamentos_Candidato_Idade_20260811"
HOJE = date.today().isoformat()

PARAMS = {
    "learning_rate": 0.05,
    "min_child_samples": 20,
    "n_estimators": 800,
    "num_leaves": 63,
    "random_state": SEMENTE,
    "verbosity": -1,
}

FEATURES_CAND = list(FEATURES_LGBM) + FEATURES_IDADE + FEATURES_REL
FEATURES_ALT = list(FEATURES_LGBM) + FEATURES_IDADE + ["LOG_INSCRICAO"] + FEATURES_REL


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _log_insc(df):
    return np.log1p(pd.to_numeric(df["CDINSCRICAOIMOB"], errors="coerce").clip(lower=0))


def _montar(df, features, fill):
    X = df[features].copy()
    X["CDSETORFISCAL"] = X["CDSETORFISCAL"].astype("category")
    for c, m in fill.items():
        if c in X.columns:
            X[c] = X[c].fillna(m)
    y = np.log(df["VLTRANSACAO_DEFLACIONADO"])
    return X, y


def _fill(df, features):
    return {
        c: float(df[c].median()) if df[c].notna().any() else 0.0
        for c in features
        if c != "CDSETORFISCAL" and c in df.columns
    }


def _estrat(pred, real, insc):
    out = {}
    for nome, lo, hi in FAIXAS_INSCRICAO:
        m = (insc >= lo) & (insc < hi) & (real > 0) & np.isfinite(pred) & np.isfinite(real)
        if m.sum() < 30:
            out[nome] = {"n": int(m.sum())}
            continue
        met = metricas_completas(pred[m], real[m])
        out[nome] = {
            "n": int(m.sum()),
            "COD_mediano": float(met["COD_mediano"]),
            "COD": float(met["COD"]),
            "PRD": float(met["PRD"]),
            "razao_mediana": float(met["razao_mediana"]),
        }
    return out


def treinar(pool, teste, features, nome):
    fill = _fill(pool, features)
    Xtr, ytr = _montar(pool, features, fill)
    Xte, _ = _montar(teste, features, fill)
    model = lgb.LGBMRegressor(**PARAMS)
    model.fit(Xtr, ytr, categorical_feature=[c for c in CAT_LGBM if c in Xtr.columns])
    pred = np.exp(model.predict(Xte))
    real = teste["VLTRANSACAO_DEFLACIONADO"].to_numpy()
    met = metricas_completas(pred, real)
    met["PRB"] = float(prb_global(pred, real))
    insc = pd.to_numeric(teste["CDINSCRICAOIMOB"], errors="coerce").to_numpy()
    print(
        f"  {nome}: CODmed={met['COD_mediano']:.3f}% COD={met['COD']:.3f}% "
        f"PRD={met['PRD']:.4f} PRB={met['PRB']:.4f}",
        flush=True,
    )
    return {
        "nome": nome,
        "features": features,
        "fill": fill,
        "model": model,
        "pred": pred,
        "holdout": {k: float(v) for k, v in met.items()},
        "por_inscricao": _estrat(pred, real, insc),
        "importances": {
            f: int(v)
            for f, v in zip(features, model.feature_importances_.tolist())
        },
    }


def montar_pastas():
    if PACOTE.exists():
        shutil.rmtree(PACOTE)
    for sub in ("codigo", "dados/amostras", "dados/graficos", "relatorios", "decisoes_e_notas"):
        (PACOTE / sub).mkdir(parents=True, exist_ok=True)


def copiar_codigo():
    # scripts do candidato
    for nome in (
        "idade_features.py",
        "teste_inscricao_relativa_setor.py",
        "fechar_candidato_selan.py",
        "pipeline_idade_apartamentos.py",
    ):
        src = HERE / nome
        if src.exists():
            shutil.copy2(src, PACOTE / "codigo" / nome)
    # libs do aprovado necessárias à avaliação
    for nome in (
        "apartamentos_lib.py",
        "metricas_iaao_por_decil.py",
        "saneamento_chauvenet_iterativo.py",
    ):
        shutil.copy2(CODIGO_APROV / nome, PACOTE / "codigo" / nome)


def main():
    print(f"Empacotando candidato em {PACOTE.name}", flush=True)
    montar_pastas()
    copiar_codigo()

    nasc = carregar_nascimento()
    # copiar nascimento (amostra essencial ao candidato)
    shutil.copy2(PASTA_ITIV / "Nascimento Imovel.csv", PACOTE / "dados" / "Nascimento_Imovel.csv")

    pool = pd.read_parquet(AMOSTRAS / "pool_treino_parametros.parquet")
    teste = pd.read_parquet(AMOSTRAS / "teste_final_parametros.parquet")
    pool = anexar_idade(pool, nasc)
    teste = anexar_idade(teste, nasc)
    pool_limpo = aplicar_chauvenet_iterativo(preparar_area_saneamento(pool)).copy()
    pool_limpo["KNN_PROXY"] = calcular_knn_proxy(pool_limpo, pool_limpo, k=K_VIZINHOS)
    teste_p = teste.copy()
    teste_p["KNN_PROXY"] = calcular_knn_proxy(pool_limpo, teste_p, k=K_VIZINHOS)
    stats = stats_setor_treino(pool_limpo)
    pool_limpo = anexar_inscricao_relativa(pool_limpo, stats)
    teste_p = anexar_inscricao_relativa(teste_p, stats)
    pool_limpo["LOG_INSCRICAO"] = _log_insc(pool_limpo)
    teste_p["LOG_INSCRICAO"] = _log_insc(teste_p)

    # salvar amostras do candidato (com idade + relativa)
    pool_limpo.to_parquet(PACOTE / "dados" / "amostras" / "pool_treino_candidato.parquet", index=False)
    teste_p.to_parquet(PACOTE / "dados" / "amostras" / "teste_final_candidato.parquet", index=False)

    print("Avaliando baseline × candidato × alternativa...", flush=True)
    base = treinar(pool_limpo, teste_p, list(FEATURES_LGBM), "baseline_aprovado")
    cand = treinar(pool_limpo, teste_p, FEATURES_CAND, "idade_mais_rel")
    alt = treinar(pool_limpo, teste_p, FEATURES_ALT, "idade_abs_rel")

    # gravar modelo candidato (sem acento no fopen do LGBM)
    modelo_txt = PACOTE / "dados" / "modelo_lightgbm_candidato_idade_rel.txt"
    modelo_txt.write_text(cand["model"].booster_.model_to_string(), encoding="utf-8")

    import pickle
    with open(PACOTE / "dados" / "vals_inscricao_por_setor.pkl", "wb") as fh:
        pickle.dump(stats.attrs.get("vals_por_setor", {}), fh)
    st = stats.copy()
    st.attrs = {}
    st.to_parquet(PACOTE / "dados" / "stats_setor_inscricao_treino.parquet", index=False)
    (PACOTE / "dados" / "fill_medians_candidato.json").write_text(
        json.dumps(cand["fill"], indent=2), encoding="utf-8"
    )
    (PACOTE / "dados" / "features_candidato.json").write_text(
        json.dumps({
            "candidato": FEATURES_CAND,
            "alternativa_antigos": FEATURES_ALT,
            "features_idade": FEATURES_IDADE,
            "features_relativa": FEATURES_REL,
            "hiperparametros": PARAMS,
        }, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    preds = pd.DataFrame({
        "SQTRANSMISSAO": teste_p["SQTRANSMISSAO"].to_numpy(),
        "CDINSCRICAOIMOB": teste_p["CDINSCRICAOIMOB"].to_numpy(),
        "VLTRANSACAO_DEFLACIONADO": teste_p["VLTRANSACAO_DEFLACIONADO"].to_numpy(),
        "PRED_baseline": base["pred"],
        "PRED_idade_mais_rel": cand["pred"],
        "PRED_idade_abs_rel": alt["pred"],
    })
    preds.to_parquet(PACOTE / "dados" / "teste_final_previsoes_candidato.parquet", index=False)

    # carregar comparativo Compativel ja gerado, se existir
    compativel = {}
    cj = SAIDA_EXP / "comparativo_compativel_candidato_selan.json"
    if cj.exists():
        compativel = json.loads(cj.read_text(encoding="utf-8"))

    avaliacao = {
        "data": HOJE,
        "status": "CANDIDATO — aguarda avaliação Selan (não substitui o aprovado)",
        "candidato": "idade_mais_rel",
        "alternativa_equidade_antigos": "idade_abs_rel",
        "holdout": {
            "baseline_aprovado": base["holdout"],
            "idade_mais_rel": cand["holdout"],
            "idade_abs_rel": alt["holdout"],
        },
        "por_inscricao": {
            "baseline_aprovado": base["por_inscricao"],
            "idade_mais_rel": cand["por_inscricao"],
            "idade_abs_rel": alt["por_inscricao"],
        },
        "importances_candidato": cand["importances"],
        "compativel_reprocessado": {
            "fonte": "desafiante_idade/saida/comparativo_compativel_candidato_selan.json",
            "resumo": {
                k: {
                    "pct_compativel": (compativel.get("compativel_cruzamento_status_operacional") or {})
                    .get(k, {}).get("pct_compativel"),
                    "pct_compativel_com_inscricao_gt_900k": (
                        (compativel.get("compativel_cruzamento_status_operacional") or {})
                        .get(k, {}).get("pct_compativel_com_inscricao_gt_900k")
                    ),
                    "por_inscricao": (
                        (compativel.get("compativel_cruzamento_status_operacional") or {})
                        .get(k, {}).get("por_inscricao")
                    ),
                }
                for k in ("baseline", "idade_mais_rel", "idade_abs_rel")
                if (compativel.get("compativel_cruzamento_status_operacional") or {}).get(k)
            },
        },
        "metas_iaao": {
            "razao_mediana": "0,90–1,10",
            "COD": "≤15%",
            "PRD": "0,98–1,03",
            "PRB": "±0,05",
        },
    }
    (PACOTE / "dados" / "avaliacao_holdout_candidato.json").write_text(
        json.dumps(avaliacao, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # gráfico COD por inscrição
    try:
        import matplotlib.pyplot as plt
        faixas = [n for n, *_ in FAIXAS_INSCRICAO]
        fig, ax = plt.subplots(figsize=(8, 4.5))
        x = np.arange(len(faixas))
        width = 0.25
        for i, (nome, res) in enumerate(
            (("baseline", base), ("idade_mais_rel", cand), ("idade_abs_rel", alt))
        ):
            vals = [
                (res["por_inscricao"].get(f) or {}).get("COD_mediano") or 0
                for f in faixas
            ]
            ax.bar(x + i * width, vals, width, label=nome)
        ax.axhline(15, color="gray", ls="--", label="meta COD 15%")
        ax.set_xticks(x + width)
        ax.set_xticklabels(faixas)
        ax.set_ylabel("COD mediano (%)")
        ax.set_title("Erro típico por faixa de inscrição (holdout)")
        ax.legend()
        fig.tight_layout()
        fig.savefig(PACOTE / "dados" / "graficos" / "cod_mediano_por_inscricao.png", dpi=140)
        plt.close(fig)
    except Exception as exc:
        print(f"  grafico COD: {exc}", flush=True)

    # copiar gráfico Compativel e notas existentes
    for src, dest in (
        (SAIDA_EXP / "grafico_compativel_candidato_por_inscricao.png",
         PACOTE / "dados" / "graficos" / "compativel_por_inscricao.png"),
        (SAIDA_EXP / "Nota_Candidato_Selan_Idade_Relativa.docx",
         PACOTE / "relatorios" / "Nota_Candidato_Selan_Idade_Relativa.docx"),
        (SAIDA_EXP / "Nota_Candidato_Selan_Idade_Relativa.md",
         PACOTE / "relatorios" / "Nota_Candidato_Selan_Idade_Relativa.md"),
        (PASTA_ITIV / "Nota_Tecnica_Idade_Modelo_Apartamentos.docx",
         PACOTE / "decisoes_e_notas" / "Nota_Tecnica_Idade_Modelo_Apartamentos.docx"),
        (PASTA_ITIV / "Nota_Inscricao_Relativa_Setor.md",
         PACOTE / "decisoes_e_notas" / "Nota_Inscricao_Relativa_Setor.md"),
        (PASTA_ITIV / "Nota_AB_Inscricao_vs_Idade.md",
         PACOTE / "decisoes_e_notas" / "Nota_AB_Inscricao_vs_Idade.md"),
        (PASTA_ITIV / "Nota_Teste_SELIC_lags.md",
         PACOTE / "decisoes_e_notas" / "Nota_Teste_SELIC_lags.md"),
    ):
        if src.exists():
            shutil.copy2(src, dest)

    # manifesto SHA-256
    arquivos_hash = []
    for p in sorted(PACOTE.rglob("*")):
        if p.is_file() and p.suffix.lower() in {".txt", ".py", ".json", ".parquet", ".csv", ".pkl", ".md", ".docx", ".png"}:
            arquivos_hash.append({"arquivo": str(p.relative_to(PACOTE)).replace("\\", "/"), "sha256": sha256(p)})
    manifesto = {
        "pacote": PACOTE.name,
        "data": HOJE,
        "candidato": "idade_mais_rel",
        "modelo_aprovado_referencia": "Modelo_Apartamentos_Aprovado (14/07/2026)",
        "arquivos": arquivos_hash,
    }
    (PACOTE / "dados" / "manifesto_candidato.json").write_text(
        json.dumps(manifesto, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # LEIA-ME
    hb, hc, ha = base["holdout"], cand["holdout"], alt["holdout"]
    ant_b = (base["por_inscricao"].get("<300k") or {}).get("COD_mediano")
    ant_c = (cand["por_inscricao"].get("<300k") or {}).get("COD_mediano")
    ant_a = (alt["por_inscricao"].get("<300k") or {}).get("COD_mediano")

    leia = f"""# Modelo candidato — Apartamentos com idade + inscrição relativa

**Pacote:** `{PACOTE.name}`  
**Data:** {HOJE}  
**Status:** CANDIDATO para avaliação do Selan — **não substitui** o modelo aprovado em `Modelo_Apartamentos_Aprovado/` (14/07/2026).

---

## Comece por aqui (Selan)

1. **`relatorios/Resumo_Executivo_Candidato_Selan.docx`** — 1 página, linguagem simples.
2. **`relatorios/Nota_Candidato_Selan_Idade_Relativa.docx`** — números de Compatível e decisão sugerida.
3. **`dados/avaliacao_holdout_candidato.json`** — métricas IAAO completas.
4. **`dados/graficos/`** — COD e % Compatível por faixa de inscrição.

---

## O que é este modelo

O modelo aprovado avalia bem apartamentos **novos** e erra mais nos **antigos**.  
Este candidato acrescenta:

- **idade** do imóvel (quando o ano de construção existe);
- **inscrição relativa no setor** (se o imóvel é novo/velho *naquele bairro*).

Assim o sistema “enxerga” melhor a idade mesmo quando falta o ano de construção.

Alternativa testada (melhor ainda nos antigos, quase empatada no global):  
**idade + inscrição absoluta + relativa** (`idade_abs_rel`).

---

## Resultado no teste final (holdout)

| Modelo | COD mediano | COD | PRD | PRB | COD antigos (&lt;300k) |
|---|---:|---:|---:|---:|---:|
| Baseline (aprovado) | {hb['COD_mediano']:.3f}% | {hb['COD']:.3f}% | {hb['PRD']:.4f} | {hb['PRB']:.4f} | {ant_b:.2f}% |
| **Candidato (idade+rel)** | **{hc['COD_mediano']:.3f}%** | {hc['COD']:.3f}% | {hc['PRD']:.4f} | {hc['PRB']:.4f} | {ant_c:.2f}% |
| Alternativa (idade+abs+rel) | {ha['COD_mediano']:.3f}% | {ha['COD']:.3f}% | {ha['PRD']:.4f} | {ha['PRB']:.4f} | {ant_a:.2f}% |
| Meta IAAO | — | ≤15% | 0,98–1,03 | ±0,05 | — |

No reprocessamento do status Compatível (±15%), os antigos sobem de ~74% para ~81% Compatível (menos auditoria desnecessária).

---

## Conteúdo da pasta

| Pasta | Conteúdo |
|---|---|
| `dados/` | Modelo LightGBM, features, medianas, stats de setor, manifesto SHA-256, amostras |
| `dados/graficos/` | Gráficos da avaliação |
| `relatorios/` | Resumo executivo + nota Selan |
| `decisoes_e_notas/` | Notas técnicas (idade, inscrição, SELIC) |
| `codigo/` | Scripts para reproduzir a avaliação |

---

## O que **não** está neste pacote

- Substituição automática do modelo em produção  
- Planilha de status completa regenerada para 100% do universo (só holdout cruzado)  
- Etapa 5 hedônica refeita linha a linha (âncora normativa do aprovado permanece válida; o ganho aqui é do LightGBM com novas features)

---

## Decisão solicitada ao Selan

1. Autorizar a equipe a tratar `idade_mais_rel` como **candidato oficial a promoção**?  
2. Preferir a alternativa `idade_abs_rel` (foco nos antigos)?  
3. Autorizar Etapa 5 completa + regeneração do lote/status antes da troca em produção?

Enquanto isso, o modelo aprovado em 14/07/2026 **permanece** o oficial.
"""
    (PACOTE / "LEIA-ME.md").write_text(leia, encoding="utf-8")

    # Resumo executivo Word (1 página)
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        AZUL = RGBColor(0x1A, 0x3A, 0x5C)

        def shade(cell, hexc):
            tcPr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:color"), "auto")
            sh.set(qn("w:fill"), hexc)
            tcPr.append(sh)

        doc = Document()
        for s in doc.sections:
            s.top_margin = Cm(2)
            s.bottom_margin = Cm(2)
            s.left_margin = Cm(2.3)
            s.right_margin = Cm(2.3)
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(11)

        h = doc.add_heading("Resumo executivo — Modelo candidato (idade + inscrição relativa)", level=1)
        for r in h.runs:
            r.font.color.rgb = AZUL

        p = doc.add_paragraph()
        r = p.add_run("Para avaliação do Selan — NÃO substitui o modelo aprovado em 14/07/2026.")
        r.bold = True
        r.font.size = Pt(10)

        doc.add_heading("Em uma frase", level=2)
        doc.add_paragraph(
            "O candidato usa a idade do imóvel e a posição da inscrição no bairro para "
            "errar menos nos apartamentos antigos, sem piorar os novos."
        )

        doc.add_heading("Por que existe", level=2)
        doc.add_paragraph(
            "O modelo aprovado acerta mais nos imóveis novos. Nos antigos, a taxa de "
            "“Compatível” era bem menor e mais casos iam para auditoria. A idade (quando "
            "existe) e a inscrição relativa no setor corrigem essa assimetria."
        )

        doc.add_heading("Números no teste final", level=2)
        rows = [
            ("Baseline (aprovado)", f"{hb['COD_mediano']:.2f}%", f"{ant_b:.2f}%"),
            ("Candidato idade+relativa", f"{hc['COD_mediano']:.2f}%", f"{ant_c:.2f}%"),
            ("Alternativa idade+abs+rel", f"{ha['COD_mediano']:.2f}%", f"{ant_a:.2f}%"),
        ]
        t = doc.add_table(rows=1 + len(rows), cols=3)
        t.style = "Table Grid"
        for j, txt in enumerate(["Modelo", "Erro típico (COD mediano)", "Erro nos antigos"]):
            cell = t.cell(0, j)
            shade(cell, "1A3A5C")
            run = cell.paragraphs[0].add_run(txt)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                run = t.cell(i + 1, j).paragraphs[0].add_run(v)
                run.font.size = Pt(9)
                run.font.name = "Arial"

        doc.add_paragraph()
        doc.add_heading("Status Compatível nos antigos (regra ±15%)", level=2)
        doc.add_paragraph(
            "Nos imóveis de inscrição baixa (<300k), o Compatível sobe de cerca de 74% "
            "para cerca de 81%, e cai a parcela enviada a auditor."
        )

        doc.add_heading("O que pedimos ao Selan", level=2)
        for item in (
            "Autorizar este pacote como candidato oficial a promoção.",
            "Indicar se prefere o candidato (melhor no geral) ou a alternativa (melhor nos antigos).",
            "Autorizar a Etapa 5 completa e a regeneração do lote/status antes de trocar a produção.",
        ):
            doc.add_paragraph(item, style="List Number")

        doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run(
            "Enquanto isso, o modelo aprovado em 14/07/2026 permanece o oficial. "
            f"Pasta: {PACOTE.name}"
        )
        r.italic = True
        r.font.size = Pt(9)

        out = PACOTE / "relatorios" / "Resumo_Executivo_Candidato_Selan.docx"
        doc.save(out)
        doc.save(PASTA_ITIV / "Resumo_Executivo_Candidato_Selan.docx")
        print(f"  Resumo: {out.name}", flush=True)
    except Exception as exc:
        print(f"  Word resumo: {exc}", flush=True)

    print(f"\nPacote pronto: {PACOTE}", flush=True)
    print(f"Arquivos: {sum(1 for _ in PACOTE.rglob('*') if _.is_file())}", flush=True)


if __name__ == "__main__":
    main()

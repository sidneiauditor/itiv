# -*- coding: utf-8 -*-
"""Gera a Nota Técnica de Idade (ITIV) em Word (.docx)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ITIV = Path(r"D:\Pai\Coordenadoria de Inteligência Fiscal\ITIV")
SAIDA = ITIV / "desafiante_idade" / "saida"


def _font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text, bold=False, size=11, after=8, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    _font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def _heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _font(run, size=16 if level == 1 else 13, bold=True)
    return p


def _shade(cell, fill="1F4E79"):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc.append(shd)


def _table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        _shade(cell)
        for p in cell.paragraphs:
            for r in p.runs:
                _font(r, size=10, bold=True, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    _font(r, size=10)
    doc.add_paragraph()
    return t


def _br(n: int | float) -> str:
    return f"{n:,.0f}".replace(",", ".")


def main() -> Path:
    diag = json.loads((SAIDA / "diagnostico_idade_status.json").read_text(encoding="utf-8"))
    treino = json.loads(
        (SAIDA / "comparativo_baseline_desafiante.json").read_text(encoding="utf-8")
    )
    david = diag["confirmacao_david"]
    base_h = treino["baseline"]["holdout"]
    chal_h = treino["desafiante"]["holdout"]
    base_est = treino["baseline"]["estratificado_holdout"]["por_inscricao"]
    chal_est = treino["desafiante"]["estratificado_holdout"]["por_inscricao"]

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    _para(
        doc,
        "Nota técnica — Idade do imóvel no modelo de apartamentos (ITIV)",
        bold=True,
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _para(
        doc,
        "Adendo analítico ao Relatório de Inteligência Fiscal 114/2026",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _para(
        doc,
        "SELAN / SEFAZ Salvador  ·  10/08/2026\n"
        "Referência: e-mails Marcos (03/08) e David (04/08/2026)\n"
        "Fonte de idade: Nascimento Imovel.csv",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    _heading(doc, "Resumo em linguagem simples", 1)
    _para(
        doc,
        "David observou que, entre os imóveis em que o modelo “bateu” com a transação "
        "(status Compatível), quase metade tem inscrição imobiliária alta (acima de 900 mil). "
        "Inscrição alta, em Salvador, costuma indicar imóvel mais novo. Em outras palavras: "
        "o modelo acerta mais nos imóveis novos do que nos antigos/depreciados.",
    )
    _para(
        doc,
        "Isso faz sentido porque o modelo oficial de apartamentos, aprovado pela equipe, "
        "não usa a idade do prédio como variável. Ele olha área, localização, andar e vizinhança. "
        "Sem a idade, o modelo “entende” melhor o estoque novo.",
    )
    _para(
        doc,
        "Com a planilha Nascimento Imóvel (ano de construção), calculamos a idade na data da venda "
        "e treinamos um modelo experimental (desafiante) incluindo essa informação — "
        "sem substituir o modelo oficial. O desafiante ficou um pouco melhor no geral e, "
        "principalmente, melhor nos imóveis antigos. A idade passou a ser uma das variáveis mais importantes.",
    )
    _para(
        doc,
        "Em uma frase: David tem razão; incluir a idade melhora o quadro, mas o modelo oficial "
        "só deve ser trocado depois de validação formal da SELAN.",
        bold=True,
    )

    _heading(doc, "1. Confirmação do achado do David", 1)
    _table(
        doc,
        ["Indicador", "Valor"],
        [
            ["N Compatível (com estimativa do modelo)", _br(diag["n_compativel_total"])],
            [
                "Compatível com inscrição > 900.000",
                f"{_br(david['n_compativel_inscricao_gt_900k'])} ({david['pct_compativel_com_inscricao_gt_900k']}%)",
            ],
            [
                "% dos Não Compatível com inscrição > 900k",
                f"{david['pct_nao_compativel_com_inscricao_gt_900k']}%",
            ],
            ["Idade mediana — Compatível", f"{david['idade_mediana_compativel']:.0f} anos"],
            [
                "Idade mediana — Não Compatível*",
                f"{david['idade_mediana_nao_compativel']:.0f} anos",
            ],
            ["Cobertura de idade na base de status", f"{david['cobertura_idade_pct']}%"],
        ],
    )
    _para(
        doc,
        "*Exceto “Fora de Escopo do Modelo”. O percentual de Compatível com inscrição > 900k "
        "(~44,5%) está alinhado à observação de David (~48%).",
        size=10,
    )

    _heading(doc, "1.1 % Compatível por faixa de inscrição", 2)
    rows_i = []
    for r in diag["por_faixa_inscricao"]:
        rows_i.append([
            r["faixa"],
            _br(r["n"]),
            f"{r['pct_compativel']}%",
            f"{r['COD_mediano']:.1f}%" if r.get("COD_mediano") is not None else "—",
            f"{r['PRD']:.3f}" if r.get("PRD") is not None else "—",
            f"{r['idade_mediana']:.0f}" if r.get("idade_mediana") is not None else "—",
        ])
    _table(
        doc,
        ["Faixa", "N", "% Compatível", "COD mediano", "PRD", "Idade mediana"],
        rows_i,
    )

    _heading(doc, "1.2 % Compatível por faixa etária", 2)
    rows_a = []
    for r in diag["por_faixa_idade"]:
        if r["faixa"] == "(sem idade)":
            continue
        rows_a.append([
            r["faixa"],
            _br(r["n"]),
            f"{r['pct_compativel']}%",
            f"{r['COD_mediano']:.1f}%" if r.get("COD_mediano") is not None else "—",
            f"{r['idade_mediana']:.0f}" if r.get("idade_mediana") is not None else "—",
        ])
    _table(
        doc,
        ["Faixa etária", "N", "% Compatível", "COD mediano", "Idade mediana"],
        rows_a,
    )

    g1 = SAIDA / "grafico_compativel_por_faixa.png"
    g2 = SAIDA / "grafico_cod_por_inscricao.png"
    if g1.exists():
        _para(doc, "Gráfico — % Compatível por faixa de inscrição e por idade:", bold=True)
        doc.add_picture(str(g1), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if g2.exists():
        _para(doc, "Gráfico — COD mediano por faixa de inscrição (modelo aprovado):", bold=True)
        doc.add_picture(str(g2), width=Cm(12))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _heading(doc, "2. O que foi feito tecnicamente", 1)
    _para(
        doc,
        "• Ligamos cada transmissão à planilha Nascimento Imóvel pela inscrição.\n"
        "• Ano de construção só entra se estiver entre 1500 e o ano da venda (higiene).\n"
        "• Idade na transação = ano da venda − ano de construção.\n"
        "• Marcamos ausência/invalidez com flags (sem “inventar” idade escondida).\n"
        "• Treinamos um LightGBM desafiante com as mesmas configurações do modelo aprovado, "
        "só acrescentando a idade — o arquivo oficial não foi alterado.",
    )

    _heading(doc, "3. Resultado do modelo desafiante", 1)
    _para(doc, "Comparação no teste final (holdout), nunca usado no treino:")
    _table(
        doc,
        ["Modelo", "Razão mediana", "COD", "COD mediano", "PRD", "PRB"],
        [
            [
                "Baseline (sem idade)",
                f"{base_h['razao_mediana']:.4f}",
                f"{base_h['COD']:.2f}%",
                f"{base_h['COD_mediano']:.2f}%",
                f"{base_h['PRD']:.4f}",
                f"{base_h['PRB']:.4f}",
            ],
            [
                "Desafiante (+ idade)",
                f"{chal_h['razao_mediana']:.4f}",
                f"{chal_h['COD']:.2f}%",
                f"{chal_h['COD_mediano']:.2f}%",
                f"{chal_h['PRD']:.4f}",
                f"{chal_h['PRB']:.4f}",
            ],
        ],
    )
    _para(
        doc,
        f"Melhora no COD mediano global: {treino['delta_holdout_cod_mediano']:.2f} ponto percentual "
        f"(de {base_h['COD_mediano']:.2f}% para {chal_h['COD_mediano']:.2f}%).",
    )
    _para(
        doc,
        "Holdout estratificado — COD mediano por inscrição (estoque antigo × novo):",
        bold=True,
    )
    rows_e = []
    for b, c in zip(base_est, chal_est):
        rows_e.append([
            b["faixa"],
            f"{b['COD_mediano']:.2f}%" if b.get("COD_mediano") is not None else "—",
            f"{c['COD_mediano']:.2f}%" if c.get("COD_mediano") is not None else "—",
        ])
    _table(doc, ["Faixa de inscrição", "Baseline", "Desafiante"], rows_e)

    imps = treino["desafiante"]["importances"][:8]
    _para(doc, "Principais variáveis do desafiante (importância no LightGBM):", bold=True)
    _para(doc, " · ".join(f"{n} ({v})" for n, v in imps))

    _heading(doc, "4. Recomendação", 1)
    _para(
        doc,
        "Manter o modelo aprovado em produção. Tratar o desafiante com idade como candidato "
        "a substituição, após revisão SELAN (holdout completo, resíduos espaciais e "
        "reprocessamento do status Compatível).",
    )
    _para(
        doc,
        "Próximos passos sugeridos:\n"
        "1. Reprocessar o status Compatível com o desafiante e repetir as tabelas do item 1.\n"
        "2. Se a equipe validar: dossiê SELAN + novo manifesto de congelamento.\n"
        "3. Se o ganho no estoque antigo for insuficiente: testar interação idade × setor fiscal.\n"
        "4. Depois de fechar apartamentos, estender o uso do nascimento a salas comerciais.",
    )

    _heading(doc, "5. Onde estão os arquivos", 1)
    _para(
        doc,
        "Pasta: ITIV\\desafiante_idade\\saida\\\n"
        "• diagnostico_idade_status.json / CSVs e gráficos\n"
        "• modelo_lightgbm_apartamentos_com_idade.txt (experimental)\n"
        "• comparativo_baseline_desafiante.json\n"
        "• Versão Markdown desta nota na mesma pasta e na raiz de ITIV.",
    )

    _para(
        doc,
        "Documento gerado automaticamente a partir dos resultados do pipeline "
        "desafiante_idade/pipeline_idade_apartamentos.py — SELAN / Coordenadoria de Inteligência Fiscal.",
        size=9,
    )

    out1 = SAIDA / "Nota_Tecnica_Idade_Modelo_Apartamentos.docx"
    out2 = ITIV / "Nota_Tecnica_Idade_Modelo_Apartamentos.docx"
    doc.save(str(out1))
    doc.save(str(out2))
    print(out1)
    print(out2)
    return out2


if __name__ == "__main__":
    main()

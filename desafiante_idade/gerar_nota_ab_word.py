# -*- coding: utf-8 -*-
"""Gera a nota A/B inscricao × idade em Word (.docx)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ITIV = Path(r"D:\Pai\Coordenadoria de Inteligência Fiscal\ITIV")
SAIDA = ITIV / "desafiante_idade" / "saida"
JSON_AB = SAIDA / "comparativo_inscricao_vs_idade.json"


def _font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text, bold=False, size=11, after=8, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    _font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def _heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _font(run, size=16 if level == 1 else 13, bold=True)
    return p


def _shade(cell, fill="1F4E79"):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc.append(shd)


def _table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        _shade(cell)
        for p in cell.paragraphs:
            for r in p.runs:
                _font(r, size=10, bold=True, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    _font(r, size=10)
    doc.add_paragraph()
    return t


def _pct(x, casas=2):
    if x is None:
        return "—"
    return f"{x:.{casas}f}%".replace(".", ",")


def _delta(x):
    return f"{x:+.3f}".replace(".", ",")


def _cod_estrat(d, chave):
    bloco = (d or {}).get(chave) or {}
    return _pct(bloco.get("COD_mediano"))


def main() -> Path:
    data = json.loads(JSON_AB.read_text(encoding="utf-8"))
    ranking = data["ranking_cod_mediano"]
    melhor = data["melhor_modelo"]
    base = data["baseline_cod_mediano"]
    cov = data.get("cobertura_pool", {})

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    _para(
        doc,
        "Nota técnica — A/B inscrição × idade no modelo de apartamentos (ITIV)",
        bold=True,
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _para(
        doc,
        "Teste da proposta de Gabriel Ramos: usar o número de inscrição "
        "como proxy de idade (cobertura plena), comparado à idade do Nascimento Imóvel",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _para(
        doc,
        "SELAN / SEFAZ Salvador  ·  10/08/2026\n"
        "Referência: mensagens Gabriel Ramos (10/08/2026) e achado de David sobre inscrição > 900 mil\n"
        "Modelo aprovado permanece congelado; este é experimento desafiante",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    _heading(doc, "Resumo em linguagem simples", 1)
    _para(
        doc,
        "Gabriel sugeriu: em vez de depender só do ano de construção (que falta em "
        "cerca de 23% dos casos), usar o número de inscrição imobiliária — que quase "
        "todos têm — como medida de “quão novo” é o imóvel.",
    )
    _para(
        doc,
        "Testamos quatro modelos no mesmo holdout: (1) o aprovado (sem idade), "
        "(2) só inscrição, (3) só idade do nascimento, (4) idade + inscrição. "
        "A inscrição sozinha já melhora o modelo; a idade sozinha melhora um pouco mais; "
        "as duas juntas ficam no topo, com ganho extra pequeno.",
    )
    _para(
        doc,
        f"Melhor candidato neste teste: «{melhor}» "
        f"(COD mediano = {_pct(data['melhor_cod_mediano'], 3)}).",
        bold=True,
    )

    _heading(doc, "Desenho do teste", 1)
    _para(
        doc,
        f"Feature de inscrição: {data['feature_inscricao']}.",
    )
    _para(
        doc,
        "Features de idade: "
        + ", ".join(data.get("features_idade", []))
        + " (fonte: Nascimento Imovel.csv).",
    )
    _para(
        doc,
        f"Cobertura no pool de treino: idade ≈ {_pct(cov.get('idade_pct'), 1)}; "
        f"log(inscrição) ≈ {_pct(cov.get('log_inscricao_pct'), 1)}.",
    )
    _para(
        doc,
        "Mesmos hiperparâmetros LightGBM do modelo aprovado; Chauvenet + KNN iguais "
        "para todas as variantes (comparação justa).",
    )

    _heading(doc, "Resultado global (holdout)", 1)
    rows = []
    rotulos = {
        "idade_mais_inscricao": "Idade + inscrição",
        "so_idade": "Só idade",
        "so_inscricao": "Só inscrição",
        "baseline_aprovado": "Baseline (aprovado)",
    }
    for r in ranking:
        rows.append(
            [
                rotulos.get(r["nome"], r["nome"]),
                _pct(r["COD_mediano"], 3),
                _delta(r["delta_vs_baseline_pp"]) + " pp",
                _pct(r["COD"], 2),
                f"{r['PRD']:.4f}".replace(".", ","),
            ]
        )
    _table(
        doc,
        ["Modelo", "COD mediano", "Δ vs baseline", "COD", "PRD"],
        rows,
    )
    _para(
        doc,
        f"Baseline COD mediano = {_pct(base, 3)}. Valores menores de COD são melhores.",
        size=10,
    )

    _heading(doc, "Por faixa de inscrição e cobertura de ano", 1)
    rows_fx = []
    for r in ranking:
        e = r.get("estratificado") or {}
        pi = e.get("por_inscricao") or {}
        pa = e.get("por_ano_cobertura") or {}
        rows_fx.append(
            [
                rotulos.get(r["nome"], r["nome"]),
                _cod_estrat(pi, "<300k"),
                _cod_estrat(pi, "300-900k"),
                _cod_estrat(pi, ">900k"),
                _cod_estrat(pa, "com_ano"),
                _cod_estrat(pa, "sem_ano"),
            ]
        )
    _table(
        doc,
        ["Modelo", "<300k", "300–900k", ">900k", "Com ano", "Sem ano"],
        rows_fx,
    )

    _heading(doc, "Interpretação", 1)
    bullets = [
        "Gabriel tem razão na direção: só a inscrição já melhora o baseline "
        "(cerca de 6,53% → 6,47% no COD mediano) e cobre praticamente 100% das amostras.",
        "A idade do nascimento ainda é um pouco melhor que só a inscrição.",
        "As duas juntas vencem por pouco: sinais quase redundantes, mas complementaridade "
        "leve (especialmente em estoque antigo e em quem não tem ano).",
        "Correlação idade × log(inscrição) no pool ≈ −0,83 "
        "(inscrição alta ↔ imóvel mais novo): proxy forte, não idêntico.",
        "Nos sem ano de construção, inscrição ajuda frente ao baseline; "
        "idade+inscrição chega ao menor erro nessa fatia.",
        "Nos antigos (inscrição < 300 mil), idade+inscrição é o melhor.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        _font(run, size=11)
        p.paragraph_format.space_after = Pt(4)

    _heading(doc, "Recomendação", 1)
    _para(
        doc,
        "Usar a inscrição (log) como feature permanente, pela cobertura total; "
        "manter a idade do nascimento quando existir. "
        "O candidato desafiante deste teste é idade + inscrição. "
        "O modelo aprovado em produção permanece congelado até validação formal da SELAN.",
        bold=True,
    )
    _para(
        doc,
        "Observação (Gabriel): o ritmo de criação de inscrição muda no tempo e pode "
        "introduzir viés de segunda ordem; fica para modelagem futura. "
        "O teste SELIC (D0/M3/M6) não mostrou ganho material neste mesmo desenho.",
        size=10,
    )

    _heading(doc, "Arquivos", 1)
    _para(
        doc,
        "desafiante_idade/saida/comparativo_inscricao_vs_idade.json\n"
        "desafiante_idade/saida/preds_inscricao_vs_idade.csv\n"
        "ITIV/Nota_AB_Inscricao_vs_Idade.md",
        size=10,
    )

    out_main = ITIV / "Nota_AB_Inscricao_vs_Idade.docx"
    out_copy = SAIDA / "Nota_AB_Inscricao_vs_Idade.docx"
    doc.save(out_main)
    doc.save(out_copy)
    print(out_main)
    print(out_copy)
    return out_main


if __name__ == "__main__":
    main()

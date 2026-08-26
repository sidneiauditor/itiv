# -*- coding: utf-8 -*-
"""Candidato Selan: idade + inscricao relativa — reprocessa status Compativel.

1) Treina e grava o candidato `idade_mais_rel` (e variante `idade_abs_rel`).
2) Gera previsoes OOF (pool) + holdout.
3) Reclassifica com a MESMA regra operacional (+/-15%, suspeito 0,5-2,0).
4) Compara % Compativel por faixa de inscricao/idade vs baseline.
5) Escreve nota para decisao Selan (md + json).

Nao altera o modelo aprovado em Modelo_Apartamentos_Aprovado/.

Executar:
  python -u desafiante_idade/reprocessar_compativel_candidato.py
"""
from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

import lightgbm as lgb
import numpy as np
import pandas as pd
from sklearn.model_selection import KFold

HERE = Path(__file__).resolve().parent
PASTA_ITIV = HERE.parent
CODIGO = PASTA_ITIV / "Modelo_Apartamentos_Aprovado" / "codigo"
sys.path.insert(0, str(CODIGO))
sys.path.insert(0, str(HERE))

from idade_features import (  # noqa: E402
    FEATURES_IDADE,
    FAIXAS_IDADE,
    FAIXAS_INSCRICAO,
    anexar_idade,
    carregar_nascimento,
    faixa_idade,
    faixa_inscricao,
)
from teste_inscricao_relativa_setor import (  # noqa: E402
    FEATURES_REL,
    anexar_inscricao_relativa,
    stats_setor_treino,
)
from apartamentos_lib import (  # noqa: E402
    AMOSTRAS,
    CAT_LGBM,
    FEATURES_LGBM,
    K_VIZINHOS,
    SEMENTE,
    calcular_knn_proxy,
    metricas_completas,
    preparar_area_saneamento,
)
from saneamento_chauvenet_iterativo import aplicar_chauvenet_iterativo  # noqa: E402

warnings.filterwarnings("ignore")

SAIDA = HERE / "saida"
SAIDA.mkdir(parents=True, exist_ok=True)

STATUS_XLSX = PASTA_ITIV / "Informações ITIV 20260610 - Status Apartamentos.xlsx"

TOLERANCIA = 0.15
LIMITE_SUSPEITO_INF = 0.5
LIMITE_SUSPEITO_SUP = 2.0
COMPATIVEL = "Compatível"
COMPATIVEL_ABAIXO = "Compatível (valor da transação abaixo do valor estimado pelo modelo)"
NECESSITA_AUDITOR = "Necessidade Avaliação por Auditor"
SUSPEITO = "Valor da Transação Suspeito (fora de padrão)"
FORA_ESCOPO = "Fora de Escopo do Modelo"

PARAMS = {
    "learning_rate": 0.05,
    "min_child_samples": 20,
    "n_estimators": 800,
    "num_leaves": 63,
    "random_state": SEMENTE,
    "verbosity": -1,
}
K_FOLDS = 5

CANDIDATOS = {
    "baseline": list(FEATURES_LGBM),
    "idade_mais_rel": list(FEATURES_LGBM) + FEATURES_IDADE + FEATURES_REL,
    "idade_abs_rel": list(FEATURES_LGBM) + FEATURES_IDADE + ["LOG_INSCRICAO"] + FEATURES_REL,
}


def classificar(valor_atualizado, estimativa) -> str:
    if pd.isna(estimativa) or estimativa <= 0:
        return FORA_ESCOPO
    if pd.isna(valor_atualizado) or valor_atualizado <= 0:
        return FORA_ESCOPO
    razao = estimativa / valor_atualizado
    if razao < LIMITE_SUSPEITO_INF or razao > LIMITE_SUSPEITO_SUP:
        return SUSPEITO
    if (1 - TOLERANCIA) <= razao <= (1 + TOLERANCIA):
        return COMPATIVEL
    if razao < (1 - TOLERANCIA):
        return COMPATIVEL_ABAIXO
    return NECESSITA_AUDITOR


def _eh_compativel(status: str) -> bool:
    return str(status or "").startswith("Compatível")


def _log_insc(df: pd.DataFrame) -> pd.Series:
    insc = pd.to_numeric(df["CDINSCRICAOIMOB"], errors="coerce")
    return np.log1p(insc.clip(lower=0))


def _montar(df, features, fill):
    X = df[features].copy()
    X["CDSETORFISCAL"] = X["CDSETORFISCAL"].astype("category")
    for col, med in fill.items():
        if col in X.columns:
            X[col] = X[col].fillna(med)
    y = np.log(df["VLTRANSACAO_DEFLACIONADO"])
    return X, y


def _fill_stats(df, features):
    fill = {}
    for c in features:
        if c == "CDSETORFISCAL":
            continue
        if c in df.columns:
            fill[c] = float(df[c].median()) if df[c].notna().any() else 0.0
    return fill


def _preparar_base(nasc: pd.DataFrame):
    pool = pd.read_parquet(AMOSTRAS / "pool_treino_parametros.parquet")
    teste = pd.read_parquet(AMOSTRAS / "teste_final_parametros.parquet")
    pool = anexar_idade(pool, nasc)
    teste = anexar_idade(teste, nasc)
    pool_san = preparar_area_saneamento(pool)
    pool_limpo = aplicar_chauvenet_iterativo(pool_san).copy()
    pool_limpo["KNN_PROXY"] = calcular_knn_proxy(pool_limpo, pool_limpo, k=K_VIZINHOS)
    teste_p = teste.copy()
    teste_p["KNN_PROXY"] = calcular_knn_proxy(pool_limpo, teste_p, k=K_VIZINHOS)
    stats = stats_setor_treino(pool_limpo)
    pool_limpo = anexar_inscricao_relativa(pool_limpo, stats)
    teste_p = anexar_inscricao_relativa(teste_p, stats)
    pool_limpo["LOG_INSCRICAO"] = _log_insc(pool_limpo)
    teste_p["LOG_INSCRICAO"] = _log_insc(teste_p)
    return pool_limpo, teste_p, stats


def treinar_e_prever(pool, teste, features, nome):
    print(f"\n=== {nome} | {len(features)} features ===", flush=True)
    fill = _fill_stats(pool, features)
    Xtr, ytr = _montar(pool, features, fill)
    Xte, _ = _montar(teste, features, fill)
    cats = [c for c in CAT_LGBM if c in Xtr.columns]

    # Holdout
    model = lgb.LGBMRegressor(**PARAMS)
    model.fit(Xtr, ytr, categorical_feature=cats)
    pred_te = np.exp(model.predict(Xte))
    real_te = teste["VLTRANSACAO_DEFLACIONADO"].to_numpy()
    met = metricas_completas(pred_te, real_te)
    print(
        f"  holdout CODmed={met['COD_mediano']:.3f}% COD={met['COD']:.3f}% "
        f"PRD={met['PRD']:.4f}",
        flush=True,
    )

    # OOF no pool (para reclassificar Compativel sem vazamento)
    oof = np.full(len(pool), np.nan)
    kf = KFold(n_splits=K_FOLDS, shuffle=True, random_state=SEMENTE)
    for i, (tr, va) in enumerate(kf.split(pool), 1):
        tr_df = pool.iloc[tr].copy()
        va_df = pool.iloc[va].copy()
        tr_df["KNN_PROXY"] = calcular_knn_proxy(tr_df, tr_df, k=K_VIZINHOS)
        va_df["KNN_PROXY"] = calcular_knn_proxy(tr_df, va_df, k=K_VIZINHOS)
        fill_f = _fill_stats(tr_df, features)
        Xtr_f, ytr_f = _montar(tr_df, features, fill_f)
        Xva_f, _ = _montar(va_df, features, fill_f)
        m = lgb.LGBMRegressor(**PARAMS)
        m.fit(Xtr_f, ytr_f, categorical_feature=cats)
        oof[va] = np.exp(m.predict(Xva_f))
        print(f"  OOF fold {i}/{K_FOLDS} ok", flush=True)

    # Modelo final no pool inteiro (artefato candidato)
    model_final = lgb.LGBMRegressor(**PARAMS)
    model_final.fit(Xtr, ytr, categorical_feature=cats)

    return {
        "nome": nome,
        "features": features,
        "holdout": {k: float(v) for k, v in met.items()},
        "pred_holdout": pred_te,
        "pred_oof_pool": oof,
        "model": model_final,
        "fill": fill,
        "importances": {
            f: int(v)
            for f, v in zip(features, model_final.feature_importances_.tolist())
            if f in (FEATURES_IDADE + ["LOG_INSCRICAO"] + FEATURES_REL + ["VAR_TENDENCIA", "LOG_AREA", "KNN_PROXY"])
        },
    }


def _tabela_compativel(df: pd.DataFrame, col_status: str, col_faixa: str, ordem: list[str]) -> list[dict]:
    rows = []
    for faixa in ordem:
        m = df[col_faixa] == faixa
        sub = df.loc[m]
        n = len(sub)
        if n == 0:
            continue
        n_c = int(sub[col_status].map(_eh_compativel).sum())
        n_aud = int((sub[col_status] == NECESSITA_AUDITOR).sum())
        rows.append({
            "faixa": faixa,
            "n": n,
            "n_compativel": n_c,
            "pct_compativel": round(100.0 * n_c / n, 2),
            "n_auditor": n_aud,
            "pct_auditor": round(100.0 * n_aud / n, 2),
            "idade_mediana": (
                float(sub["IDADE_NA_TRANSACAO"].median())
                if sub["IDADE_NA_TRANSACAO"].notna().any() else None
            ),
        })
    return rows


def carregar_status_enxuto() -> pd.DataFrame:
    preview = pd.read_excel(STATUS_XLSX, sheet_name="Exportar Planilha", nrows=0)
    wanted, colmap = [], {}
    for c in preview.columns:
        cu = str(c).upper()
        cu_ascii = " ".join("".join(ch if ord(ch) < 128 else " " for ch in cu).split())
        dest = None
        if str(c).strip().upper() == "CDINSCRICAOIMOB":
            dest = "CDINSCRICAOIMOB"
        elif str(c).strip().upper() == "SQTRANSMISSAO":
            dest = "SQTRANSMISSAO"
        elif "ESTIMATIVA DO MODELO" in cu and "LIGHTGBM" in cu:
            dest = "VALOR_MODELO"
        elif "STATUS DA AVAL" in cu or "STATUS DA AVAL" in cu_ascii:
            dest = "STATUS_ORIGINAL"
        elif ("VALOR DA TRANS" in cu or "VALOR DA TRANS" in cu_ascii) and "ATUALIZADO" in cu_ascii:
            dest = "VALOR_TX"
        if dest:
            wanted.append(c)
            colmap[c] = dest
    print(f"  Status: lendo {[colmap[c] for c in wanted]}", flush=True)
    bruto = pd.read_excel(STATUS_XLSX, sheet_name="Exportar Planilha", usecols=wanted)
    df = bruto.rename(columns=colmap)
    df["VALOR_MODELO"] = pd.to_numeric(df["VALOR_MODELO"], errors="coerce")
    df["VALOR_TX"] = pd.to_numeric(df["VALOR_TX"], errors="coerce")
    df["CDINSCRICAOIMOB"] = pd.to_numeric(df["CDINSCRICAOIMOB"], errors="coerce")
    df = df[df["VALOR_MODELO"].notna() & (df["VALOR_MODELO"] > 0)].copy()
    return df


def main():
    print("Candidato Selan — reprocessar Compativel", flush=True)
    nasc = carregar_nascimento()
    pool, teste, stats = _preparar_base(nasc)
    print(f"  pool={len(pool):,} | teste={len(teste):,} | setores={len(stats)}", flush=True)

    resultados = {}
    for nome, feats in CANDIDATOS.items():
        resultados[nome] = treinar_e_prever(pool, teste, feats, nome)

    # Persistir candidato principal (idade + relativa)
    # LightGBM fopen no Windows quebra com acentos no path — grava via Python.
    cand = resultados["idade_mais_rel"]
    modelo_path = SAIDA / "modelo_lightgbm_candidato_idade_rel.txt"
    modelo_path.write_text(cand["model"].booster_.model_to_string(), encoding="utf-8")
    # attrs contem ndarrays (vals_por_setor) — nao serializam em parquet
    import pickle

    vals_map = stats.attrs.get("vals_por_setor", {})
    with open(SAIDA / "vals_inscricao_por_setor.pkl", "wb") as fh:
        pickle.dump(vals_map, fh)
    stats_save = stats.copy()
    stats_save.attrs = {}
    stats_save.to_parquet(SAIDA / "stats_setor_inscricao_treino.parquet", index=False)
    print(f"  Modelo candidato salvo: {modelo_path.name}", flush=True)

    # Painel de previsoes (pool OOF + holdout)
    painel_pool = pd.DataFrame({
        "SQTRANSMISSAO": pool["SQTRANSMISSAO"].to_numpy(),
        "CDINSCRICAOIMOB": pool["CDINSCRICAOIMOB"].to_numpy(),
        "IDADE_NA_TRANSACAO": pool["IDADE_NA_TRANSACAO"].to_numpy(),
        "VLTRANSACAO_DEFLACIONADO": pool["VLTRANSACAO_DEFLACIONADO"].to_numpy(),
        "origem": "pool_oof",
    })
    painel_te = pd.DataFrame({
        "SQTRANSMISSAO": teste["SQTRANSMISSAO"].to_numpy(),
        "CDINSCRICAOIMOB": teste["CDINSCRICAOIMOB"].to_numpy(),
        "IDADE_NA_TRANSACAO": teste["IDADE_NA_TRANSACAO"].to_numpy(),
        "VLTRANSACAO_DEFLACIONADO": teste["VLTRANSACAO_DEFLACIONADO"].to_numpy(),
        "origem": "holdout",
    })
    for nome, r in resultados.items():
        painel_pool[f"PRED_{nome}"] = r["pred_oof_pool"]
        painel_te[f"PRED_{nome}"] = r["pred_holdout"]
    painel = pd.concat([painel_pool, painel_te], ignore_index=True)
    painel["FAIXA_INSCRICAO"] = faixa_inscricao(painel["CDINSCRICAOIMOB"])
    painel["FAIXA_IDADE"] = faixa_idade(painel["IDADE_NA_TRANSACAO"])

    # Classificacao com valor deflacionado da amostra (holdout+OOF)
    for nome in CANDIDATOS:
        painel[f"STATUS_{nome}"] = [
            classificar(v, e)
            for v, e in zip(painel["VLTRANSACAO_DEFLACIONADO"], painel[f"PRED_{nome}"])
        ]

    # Cruzamento com status operacional (mesmo VALOR_TX da planilha David)
    status = carregar_status_enxuto()
    # Sem data na leitura enxuta: idade usa ano de referencia cadastral (diagnostico)
    status = anexar_idade(status, nasc)

    cruz = status.merge(
        painel[["SQTRANSMISSAO"] + [f"PRED_{n}" for n in CANDIDATOS]],
        on="SQTRANSMISSAO",
        how="inner",
    )
    print(f"  Cruzamento status × amostra: {len(cruz):,} linhas", flush=True)

    for nome in CANDIDATOS:
        cruz[f"STATUS_NOVO_{nome}"] = [
            classificar(v, e) for v, e in zip(cruz["VALOR_TX"], cruz[f"PRED_{nome}"])
        ]
    # Baseline operacional = status original da planilha (modelo aprovado no lote)
    cruz["STATUS_NOVO_aprovado_planilha"] = cruz["STATUS_ORIGINAL"]

    cruz["FAIXA_INSCRICAO"] = faixa_inscricao(cruz["CDINSCRICAOIMOB"])
    cruz["FAIXA_IDADE"] = faixa_idade(cruz["IDADE_NA_TRANSACAO"])

    ordem_i = [n for n, *_ in FAIXAS_INSCRICAO]
    ordem_a = [n for n, *_ in FAIXAS_IDADE] + ["(sem idade)"]

    def resumo_modelo(df, col_status, rotulo):
        n = len(df)
        n_c = int(df[col_status].map(_eh_compativel).sum())
        n_c900 = int(
            (
                df[col_status].map(_eh_compativel)
                & (pd.to_numeric(df["CDINSCRICAOIMOB"], errors="coerce") > 900_000)
            ).sum()
        )
        return {
            "rotulo": rotulo,
            "n": n,
            "n_compativel": n_c,
            "pct_compativel": round(100.0 * n_c / n, 2) if n else None,
            "pct_compativel_com_inscricao_gt_900k": (
                round(100.0 * n_c900 / n_c, 2) if n_c else None
            ),
            "por_inscricao": _tabela_compativel(df, col_status, "FAIXA_INSCRICAO", ordem_i),
            "por_idade": _tabela_compativel(df, col_status, "FAIXA_IDADE", ordem_a),
            "contagem_status": df[col_status].astype(str).value_counts().to_dict(),
        }

    # A) Amostra (OOF+holdout) com valor deflacionado
    amostra_status = {
        nome: resumo_modelo(painel, f"STATUS_{nome}", nome) for nome in CANDIDATOS
    }

    # B) Universo do cruzamento com VALOR_TX operacional
    operacional = {
        "aprovado_planilha": resumo_modelo(cruz, "STATUS_ORIGINAL", "aprovado_planilha"),
    }
    for nome in CANDIDATOS:
        operacional[nome] = resumo_modelo(cruz, f"STATUS_NOVO_{nome}", nome)

    # Holdout puro (metricas IAAO ja conhecidas + Compativel)
    hold = painel[painel["origem"] == "holdout"].copy()
    holdout_comp = {nome: resumo_modelo(hold, f"STATUS_{nome}", nome) for nome in CANDIDATOS}

    pacote = {
        "candidato_recomendado": "idade_mais_rel",
        "alternativa_equidade_antigos": "idade_abs_rel",
        "regra_status": {
            "tolerancia": TOLERANCIA,
            "suspeito": [LIMITE_SUSPEITO_INF, LIMITE_SUSPEITO_SUP],
            "fonte": "gerar_status_avaliacao_raw.classificar (17/07/2026)",
        },
        "holdout_metricas": {n: resultados[n]["holdout"] for n in CANDIDATOS},
        "importances_foco": {n: resultados[n]["importances"] for n in CANDIDATOS},
        "compativel_holdout": holdout_comp,
        "compativel_amostra_oof_holdout": amostra_status,
        "compativel_cruzamento_status_operacional": operacional,
        "n_cruzamento": int(len(cruz)),
        "artefatos": {
            "modelo": str(modelo_path.name),
            "stats_setor": "stats_setor_inscricao_treino.parquet",
            "preds": "preds_candidato_selan.parquet",
        },
    }

    painel.to_parquet(SAIDA / "preds_candidato_selan.parquet", index=False)
    cruz.to_parquet(SAIDA / "cruzamento_status_candidato.parquet", index=False)
    out_json = SAIDA / "comparativo_compativel_candidato_selan.json"
    out_json.write_text(json.dumps(pacote, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    # Graficos
    try:
        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(8, 4.5))
        x = np.arange(len(ordem_i))
        width = 0.25
        for i, nome in enumerate(["baseline", "idade_mais_rel", "idade_abs_rel"]):
            tab = {r["faixa"]: r["pct_compativel"] for r in operacional[nome]["por_inscricao"]}
            vals = [tab.get(f, 0) for f in ordem_i]
            ax.bar(x + i * width, vals, width, label=nome)
        ax.set_xticks(x + width)
        ax.set_xticklabels(ordem_i)
        ax.set_ylabel("% Compatível")
        ax.set_ylim(0, 100)
        ax.set_title("% Compatível por inscrição (cruzamento status × amostra)")
        ax.legend()
        fig.tight_layout()
        fig.savefig(SAIDA / "grafico_compativel_candidato_por_inscricao.png", dpi=140)
        plt.close(fig)
    except Exception as exc:
        print(f"  (grafico: {exc})", flush=True)

    # Nota Selan
    def _fmt_tab(rows):
        linhas = ["| Faixa | N | % Compatível | % Auditor |", "|---|---:|---:|---:|"]
        for r in rows:
            linhas.append(
                f"| {r['faixa']} | {r['n']:,} | {r['pct_compativel']:.1f}% | {r['pct_auditor']:.1f}% |"
            )
        return "\n".join(linhas)

    b = operacional["baseline"]
    r = operacional["idade_mais_rel"]
    a = operacional["idade_abs_rel"]
    h = pacote["holdout_metricas"]

    nota = f"""# Nota para decisão Selan — candidato idade + inscrição relativa

**Data:** 11/08/2026  
**Status:** experimental — **não** substitui o modelo aprovado (14/07/2026) até decisão formal.  
**Candidato recomendado:** `idade_mais_rel` (features do aprovado + idade + inscrição relativa no setor).  
**Alternativa (melhor nos antigos):** `idade_abs_rel` (idade + inscrição absoluta + relativa).

---

## 1. O que foi feito

1. Congelou-se o candidato **idade + inscrição relativa no setor** (`LOG_INSCRICAO_REL`, `RANK_INSCRICAO_SETOR`).
2. Treinou-se com os mesmos hiperparâmetros do LightGBM aprovado.
3. Reprocessou-se o status **Compatível** com a **mesma regra operacional** (±15%; suspeito fora de 0,5–2,0).
4. Comparou-se baseline × candidatos no holdout e no cruzamento com a planilha de status.

SELIC permanece **fora** (teste anterior sem ganho).

---

## 2. Métricas no holdout (nunca usado no treino)

| Modelo | COD mediano | COD | PRD |
|---|---:|---:|---:|
| baseline | {h['baseline']['COD_mediano']:.3f}% | {h['baseline']['COD']:.3f}% | {h['baseline']['PRD']:.4f} |
| idade_mais_rel | {h['idade_mais_rel']['COD_mediano']:.3f}% | {h['idade_mais_rel']['COD']:.3f}% | {h['idade_mais_rel']['PRD']:.4f} |
| idade_abs_rel | {h['idade_abs_rel']['COD_mediano']:.3f}% | {h['idade_abs_rel']['COD']:.3f}% | {h['idade_abs_rel']['PRD']:.4f} |

---

## 3. % Compatível por inscrição (cruzamento status operacional × amostra)

Universo: {len(cruz):,} transações presentes na planilha de status **e** na amostra (OOF/holdout).  
Valor de referência: `VALOR_TX` da planilha (IPCA). Estimativa: previsão do respectivo modelo.

### Baseline (features do aprovado, reclassificado)
{_fmt_tab(b['por_inscricao'])}

### Candidato idade + relativa
{_fmt_tab(r['por_inscricao'])}

### Alternativa idade + abs + rel
{_fmt_tab(a['por_inscricao'])}

**Leitura David:** % dos Compatível com inscrição > 900k  
- baseline reclassificado: {b.get('pct_compativel_com_inscricao_gt_900k')}%  
- idade_mais_rel: {r.get('pct_compativel_com_inscricao_gt_900k')}%  
- idade_abs_rel: {a.get('pct_compativel_com_inscricao_gt_900k')}%

(Quanto menor essa concentração nos novos, mais equilibrada a taxa Compatível.)

---

## 4. % Compatível por idade (candidato idade_mais_rel)

{_fmt_tab(r['por_idade'])}

---

## 5. Recomendação à equipe / Selan

| Pergunta | Sugestão |
|---|---|
| Promover agora a produção? | **Não automaticamente** — falta Etapa 5 normativa completa + manifesto SHA-256 do candidato |
| Qual candidato preferir? | **idade_mais_rel** (melhor COD global nos testes) |
| Se prioridade = equidade nos antigos? | Considerar **idade_abs_rel** |
| O que fazer com o aprovado? | Mantê-lo congelado até aprovação formal |

### Checklist para promoção

- [ ] Revisar esta nota e os gráficos em `desafiante_idade/saida/`
- [ ] Rodar Etapa 5 normativa no candidato (hedônico + razões por decil/setor)
- [ ] Regenerar lote + planilha de status completa com o candidato
- [ ] Gerar manifesto SHA-256
- [ ] Decisão formal no `REGISTRO_ETAPAS_APROVACAO`

---

## 6. Artefatos

| Arquivo | Conteúdo |
|---|---|
| `modelo_lightgbm_candidato_idade_rel.txt` | LightGBM candidato |
| `stats_setor_inscricao_treino.parquet` | Medianas/ranks por setor (treino) |
| `preds_candidato_selan.parquet` | Previsões OOF + holdout |
| `cruzamento_status_candidato.parquet` | Status reprocessado no cruzamento |
| `comparativo_compativel_candidato_selan.json` | Pacote numérico completo |
| `grafico_compativel_candidato_por_inscricao.png` | Comparativo visual |

---

*Coordenadoria de Inteligência Fiscal — SEFAZ Salvador*
"""
    (SAIDA / "Nota_Candidato_Selan_Idade_Relativa.md").write_text(nota, encoding="utf-8")
    (PASTA_ITIV / "Nota_Candidato_Selan_Idade_Relativa.md").write_text(nota, encoding="utf-8")
    print(f"\nNota: Nota_Candidato_Selan_Idade_Relativa.md", flush=True)
    print(f"JSON: {out_json}", flush=True)

    # Word
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        import re

        AZUL = RGBColor(0x1A, 0x3A, 0x5C)

        def shade(cell, hexc):
            tcPr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:color"), "auto")
            sh.set(qn("w:fill"), hexc)
            tcPr.append(sh)

        doc = Document()
        for s in doc.sections:
            s.top_margin = Cm(2.2)
            s.bottom_margin = Cm(2.2)
            s.left_margin = Cm(2.5)
            s.right_margin = Cm(2.2)
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(11)

        for ln in nota.split("\n"):
            s = ln.rstrip()
            if not s.strip():
                continue
            if s.startswith("# "):
                p = doc.add_heading(s[2:], level=1)
                for r_ in p.runs:
                    r_.font.color.rgb = AZUL
            elif s.startswith("## "):
                p = doc.add_heading(s[3:], level=2)
                for r_ in p.runs:
                    r_.font.color.rgb = AZUL
            elif s.startswith("### "):
                p = doc.add_heading(s[4:], level=3)
                for r_ in p.runs:
                    r_.font.color.rgb = AZUL
            elif s.startswith("|") and "---" not in s:
                # acumula tabelas simples via paragrafos se quebrado — skip sep
                continue
            elif s.startswith("|"):
                continue
            elif s.startswith("- ["):
                doc.add_paragraph(s[2:], style="List Bullet")
            elif s.startswith("- "):
                doc.add_paragraph(s[2:], style="List Bullet")
            elif re.match(r"^\d+\. ", s):
                doc.add_paragraph(re.sub(r"^\d+\. ", "", s), style="List Number")
            elif s.startswith("---"):
                continue
            else:
                p = doc.add_paragraph()
                parts = re.split(r"(\*\*.+?\*\*)", s)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        run = p.add_run(part)
                    run.font.name = "Arial"

        # Tabelas principais no Word
        doc.add_heading("Anexo — Compatível por inscrição (candidato)", level=2)
        for titulo, rows in (
            ("Baseline", b["por_inscricao"]),
            ("idade_mais_rel", r["por_inscricao"]),
            ("idade_abs_rel", a["por_inscricao"]),
        ):
            doc.add_heading(titulo, level=3)
            t = doc.add_table(rows=1 + len(rows), cols=4)
            t.style = "Table Grid"
            hdr = ["Faixa", "N", "% Compatível", "% Auditor"]
            for j, hcell in enumerate(hdr):
                cell = t.cell(0, j)
                shade(cell, "1A3A5C")
                run = cell.paragraphs[0].add_run(hcell)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)
                run.font.name = "Arial"
            for i, row in enumerate(rows):
                vals = [
                    row["faixa"],
                    f"{row['n']:,}",
                    f"{row['pct_compativel']:.1f}%",
                    f"{row['pct_auditor']:.1f}%",
                ]
                for j, v in enumerate(vals):
                    cell = t.cell(i + 1, j)
                    run = cell.paragraphs[0].add_run(v)
                    run.font.size = Pt(9)
                    run.font.name = "Arial"
            doc.add_paragraph()

        docx_path = SAIDA / "Nota_Candidato_Selan_Idade_Relativa.docx"
        doc.save(docx_path)
        doc.save(PASTA_ITIV / "Nota_Candidato_Selan_Idade_Relativa.docx")
        print(f"Word: {docx_path.name}", flush=True)
    except Exception as exc:
        print(f"  (Word nao gerado: {exc})", flush=True)

    print("\n=== RESUMO Compativel (cruzamento operacional) ===", flush=True)
    for nome in ["baseline", "idade_mais_rel", "idade_abs_rel"]:
        o = operacional[nome]
        print(
            f"  {nome}: Compat={o['pct_compativel']}% | "
            f"%Compat>900k={o['pct_compativel_com_inscricao_gt_900k']}%",
            flush=True,
        )
        for row in o["por_inscricao"]:
            print(
                f"    {row['faixa']}: Compat={row['pct_compativel']}% "
                f"Auditor={row['pct_auditor']}% n={row['n']}",
                flush=True,
            )


if __name__ == "__main__":
    main()

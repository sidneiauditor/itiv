# -*- coding: utf-8 -*-
"""Converte os documentos .md do projeto para Word, e gera o doc dos 20 exemplos."""
import re
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1a, 0x3a, 0x5c)
AZ_HEX = "1A3A5C"
CIN_HEX = "F2F2F2"

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc)
    tcPr.append(sh)

def add_runs(par, texto):
    # negrito **...**
    partes = re.split(r"(\*\*.+?\*\*)", texto)
    for p in partes:
        if p.startswith("**") and p.endswith("**"):
            r = par.add_run(p[2:-2]); r.font.bold = True
        else:
            r = par.add_run(p)
        r.font.name = "Arial"

def montar_tabela(doc, linhas):
    linhas = [l for l in linhas if not re.match(r"^\|[\s\-:|]+\|$", l)]
    matriz = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas]
    if not matriz: return
    ncol = len(matriz[0])
    t = doc.add_table(rows=len(matriz), cols=ncol)
    t.style = "Table Grid"
    for i, row in enumerate(matriz):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            if i == 0:
                shade(cell, AZ_HEX)
                r = par.add_run(txt); r.font.bold = True
                r.font.color.rgb = RGBColor(0xff,0xff,0xff); r.font.size = Pt(9.5); r.font.name="Arial"
            else:
                shade(cell, CIN_HEX if i % 2 == 0 else "FFFFFF")
                add_runs(par, txt)
                for rr in par.runs: rr.font.size = Pt(9.5)
    doc.add_paragraph()

def converter(md_path, docx_path, titulo_capa=None):
    with open(md_path, encoding="utf-8") as f:
        linhas = f.read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.5); s.right_margin = Cm(2.2)
    # rodape
    fp = doc.sections[0].footer.paragraphs[0]
    rr = fp.add_run("Coordenadoria de Inteligência Fiscal — SEFAZ Salvador | Junho 2026")
    rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(0x88,0x88,0x88); rr.font.name="Arial"

    buf_tab = []
    for ln in linhas:
        if ln.strip().startswith("|"):
            buf_tab.append(ln); continue
        if buf_tab:
            montar_tabela(doc, buf_tab); buf_tab = []
        s = ln.rstrip()
        if not s.strip():
            continue
        if s.startswith("### "):
            p = doc.add_heading(s[4:], level=3)
            for r in p.runs: r.font.color.rgb = AZUL; r.font.name="Arial"
        elif s.startswith("## "):
            p = doc.add_heading(s[3:], level=2)
            for r in p.runs: r.font.color.rgb = AZUL; r.font.name="Arial"
        elif s.startswith("# "):
            p = doc.add_heading(s[2:], level=1)
            for r in p.runs: r.font.color.rgb = AZUL; r.font.name="Arial"
        elif s.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            r = p.add_run(s[2:].replace("**","")); r.font.italic = True
            r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x44,0x44,0x44); r.font.name="Arial"
        elif re.match(r"^[-*] ", s):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ","",s))
        elif s.startswith("---") or s.startswith("***"):
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(); add_runs(p, s)
    if buf_tab: montar_tabela(doc, buf_tab)
    doc.save(docx_path)
    print("OK:", docx_path)

RAIZ = r"D:\Pai\Coordenadoria de Inteligência Fiscal\ITIV"
docs = [
    ("PROJETO_Conformidade_Normas_Avaliacao.md", "PROJETO_Conformidade_Normas_Avaliacao.docx"),
    ("FUNDAMENTACAO_NORMATIVA.md", "FUNDAMENTACAO_NORMATIVA.docx"),
    ("REGISTRO_ETAPAS_APROVACAO.md", "REGISTRO_ETAPAS_APROVACAO.docx"),
]
for md, dx in docs:
    converter(f"{RAIZ}\\{md}", f"{RAIZ}\\{dx}")

# ---- Documento dos 20 exemplos de duplicatas ----
df = pd.read_parquet(f"{RAIZ}\\base_limpa_normas.parquet")
chave = ["CDINSCRICAOIMOB","DATA_TRANSACAO","VLTRANSACAO"]
dup = df[df.duplicated(subset=chave, keep=False)].sort_values(chave)
cols = ["SQTRANSMISSAO","CDINSCRICAOIMOB","DATA_TRANSACAO","VLTRANSACAO","DSSUBUNIDADE","VLAREAUSOPRIV","NUSUBUNIDADE"]
ex = dup[cols].head(20).copy()
ex["DATA_TRANSACAO"] = pd.to_datetime(ex["DATA_TRANSACAO"]).dt.strftime("%d/%m/%Y")
ex["VLTRANSACAO"] = ex["VLTRANSACAO"].map(lambda v: f"{v:,.0f}".replace(",","."))

doc = Document()
st = doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(11)
for s in doc.sections:
    s.left_margin=Cm(2); s.right_margin=Cm(2)
p = doc.add_heading("Exemplos de Transações Duplicadas — Para Aprovação da Equipe", level=1)
for r in p.runs: r.font.color.rgb=AZUL
doc.add_paragraph(
    "A validação de consistência (NBR 14653-2) identificou 13.239 linhas que compartilham "
    "a MESMA inscrição imobiliária + MESMA data + MESMO valor, distribuídas em 4.586 grupos. "
    "Abaixo, 20 exemplos reais. NÃO se conclui a causa — a equipe deve avaliar se são "
    "duplicatas a remover, vendas conjuntas (ex.: imóvel + garagem), ou reentradas."
)
hdr = ["Nº Transmissão","Inscrição","Data","Valor (R$)","Tipologia","Área m²","Subunid."]
t = doc.add_table(rows=1+len(ex), cols=len(hdr)); t.style="Table Grid"
for j,h in enumerate(hdr):
    c=t.cell(0,j); shade(c,AZ_HEX)
    r=c.paragraphs[0].add_run(h); r.font.bold=True; r.font.color.rgb=RGBColor(0xff,0xff,0xff)
    r.font.size=Pt(9); r.font.name="Arial"
for i,(_,row) in enumerate(ex.iterrows()):
    vals=[str(row["SQTRANSMISSAO"]),str(int(row["CDINSCRICAOIMOB"])),row["DATA_TRANSACAO"],
          row["VLTRANSACAO"],str(row["DSSUBUNIDADE"]),
          (str(int(row["VLAREAUSOPRIV"])) if pd.notna(row["VLAREAUSOPRIV"]) else "—"),
          (str(row["NUSUBUNIDADE"]) if pd.notna(row["NUSUBUNIDADE"]) else "—")]
    for j,v in enumerate(vals):
        c=t.cell(i+1,j); shade(c, CIN_HEX if i%2==0 else "FFFFFF")
        r=c.paragraphs[0].add_run(v); r.font.size=Pt(9); r.font.name="Arial"
doc.add_paragraph()
p=doc.add_paragraph()
r=p.add_run("Observação: nos exemplos, vários grupos têm número de transmissão DIFERENTE "
    "mas inscrição, data, valor, área e subunidade IDÊNTICOS (ex.: inscrição 10).")
r.font.size=Pt(10); r.font.italic=True
doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run("Decisão da equipe:  (  ) Remover duplicatas reais   "
    "(  ) Investigar mais   (  ) Manter (são vendas conjuntas)")
r.font.bold=True
doc.save(f"{RAIZ}\\Exemplos_Duplicatas_Aprovacao.docx")
print("OK: Exemplos_Duplicatas_Aprovacao.docx")
